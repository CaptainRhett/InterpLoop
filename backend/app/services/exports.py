"""Shared CSV/XLSX downloads for learning records."""

import csv
import io
from datetime import datetime

from flask import abort, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def export_table(headers, rows, filename, sheet_name):
    extension = "xlsx" if request.path.endswith(".xlsx") else "csv"
    if extension == "csv":
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow(headers)
        for row in rows:
            # Treat user-entered text as text when opened in spreadsheet apps.
            writer.writerow([
                "'" + value if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@"))
                else value for value in row
            ])
        data = io.BytesIO(output.getvalue().encode("utf-8-sig"))
        mimetype = "text/csv"
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        sheet.append(headers)
        for row in rows:
            if any(isinstance(value, str) and len(value) > 32767 for value in row):
                abort(400, "内容超过 Excel 单元格的 32767 字符限制，请使用 CSV 格式导出完整记录")
            sheet.append(row)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                    cell.number_format = "@"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="183A56")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 30
        for index, title in enumerate(headers, 1):
            width = 60 if title in {"源语", "ASR识别文本", "AI评价", "参考译法", "原文", "优点", "问题", "建议", "总评"} else 22
            sheet.column_dimensions[get_column_letter(index)].width = width
        data = io.BytesIO()
        workbook.save(data)
        workbook.close()
        data.seek(0)
        mimetype = XLSX_MIME
    return send_file(
        data, mimetype=mimetype, as_attachment=True,
        download_name=f"{filename}-{datetime.now().strftime('%Y%m%d')}.{extension}",
    )
