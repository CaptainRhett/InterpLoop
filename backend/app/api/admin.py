import csv
import io
import secrets
import string

from flask import Blueprint, abort, current_app, jsonify, request

from ..audit import add_audit
from ..models import (
    AcademicTerm,
    AuditLog,
    ClassGroup,
    Course,
    CourseEnrollment,
    TeachingAssignment,
    User,
    UserAccount,
    db,
)
from .auth import require_admin

admin_bp = Blueprint("admin", __name__)

HEADER_ALIASES = {
    "student_no": {"学号", "student_no", "studentid", "student_id", "账号", "login_id"},
    "name": {"姓名", "name", "学生姓名"},
    "class_code": {"班级", "班级代码", "class", "class_code"},
    "class_name": {"班级名称", "class_name"},
    "course_code": {"课程", "课程代码", "course", "course_code"},
    "course_name": {"课程名称", "course_name"},
    "term_code": {"学期", "学期代码", "term", "term_code", "semester"},
    "term_name": {"学期名称", "term_name", "semester_name"},
    "initial_password": {"初始密码", "initial_password", "password"},
}


def _json():
    return request.get_json(silent=True) or {}


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _normalize_header(value):
    return _text(value).lower().replace(" ", "")


def _canonical_row(raw):
    normalized = {_normalize_header(key): _text(value) for key, value in raw.items()}
    result = {}
    for canonical, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            key = _normalize_header(alias)
            if key in normalized and normalized[key]:
                result[canonical] = normalized[key]
                break
        result.setdefault(canonical, "")
    result["class_name"] = result["class_name"] or result["class_code"]
    result["course_name"] = result["course_name"] or result["course_code"]
    result["term_name"] = result["term_name"] or result["term_code"]
    return result


def _read_import_rows(upload):
    filename = (upload.filename or "").lower()
    content = upload.read()
    if filename.endswith(".csv"):
        text = content.decode("utf-8-sig")
        return [_canonical_row(row) for row in csv.DictReader(io.StringIO(text))]
    if filename.endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = next(values, None)
        if not headers:
            return []
        return [
            _canonical_row(dict(zip(headers, row)))
            for row in values
            if any(value is not None and _text(value) for value in row)
        ]
    abort(400, "仅支持 .xlsx 或 .csv 名单文件")


def _temporary_password(length=12):
    alphabet = string.ascii_letters + string.digits
    while True:
        password = "".join(secrets.choice(alphabet) for _ in range(length))
        if any(char.islower() for char in password) and any(
            char.isupper() for char in password
        ) and any(char.isdigit() for char in password):
            return password


def _validate_password(password):
    if len(password) < current_app.config["PASSWORD_MIN_LENGTH"]:
        abort(400, f"密码至少需要 {current_app.config['PASSWORD_MIN_LENGTH']} 位")


def _academic_items(term_id, class_id, course_id):
    term = db.session.get(AcademicTerm, term_id)
    class_group = db.session.get(ClassGroup, class_id)
    course = db.session.get(Course, course_id)
    if not term or not class_group or not course:
        abort(400, "学期、班级或课程不存在")
    if class_group.term_id != term.id or course.term_id != term.id:
        abort(400, "班级和课程必须属于同一学期")
    return term, class_group, course


@admin_bp.get("/overview")
def overview():
    require_admin()
    return jsonify(
        {
            "users": [row.to_dict() for row in User.query.order_by(User.id.desc()).limit(500)],
            "terms": [row.to_dict() for row in AcademicTerm.query.order_by(AcademicTerm.id.desc())],
            "classes": [row.to_dict() for row in ClassGroup.query.order_by(ClassGroup.id.desc())],
            "courses": [row.to_dict() for row in Course.query.order_by(Course.id.desc())],
            "enrollments": [
                row.to_dict() for row in CourseEnrollment.query.order_by(CourseEnrollment.id.desc())
            ],
            "assignments": [
                row.to_dict()
                for row in TeachingAssignment.query.order_by(TeachingAssignment.id.desc())
            ],
        }
    )


@admin_bp.get("/audit-logs")
def audit_logs():
    require_admin()
    rows = AuditLog.query.order_by(AuditLog.created_at.desc()).limit(500).all()
    return jsonify({"audit_logs": [row.to_dict() for row in rows]})


@admin_bp.post("/users")
def create_user():
    actor = require_admin()
    data = _json()
    role = data.get("role") or "student"
    if role not in {"student", "teacher", "admin"}:
        abort(400, "账号角色无效")
    login_id = (data.get("login_id") or "").strip()
    name = (data.get("name") or "").strip()
    student_no = (data.get("student_no") or "").strip() or None
    if not login_id or not name:
        abort(400, "登录账号和姓名不能为空")
    if role == "student" and not student_no:
        student_no = login_id
    if UserAccount.query.filter_by(login_id=login_id).first():
        abort(400, "登录账号已存在")
    if student_no and User.query.filter_by(student_no=student_no).first():
        abort(400, "学号或人员编号已存在")

    password = data.get("password") or _temporary_password()
    _validate_password(password)
    user = User(student_no=student_no, name=name, role=role)
    db.session.add(user)
    db.session.flush()
    account = UserAccount(user_id=user.id, login_id=login_id)
    account.set_password(password, must_change=True)
    db.session.add(account)
    add_audit(actor, "admin.user_created", "user", user.id, {"role": role})
    db.session.commit()
    return jsonify({"user": user.to_dict(), "initial_password": password}), 201


@admin_bp.patch("/users/<int:user_id>")
def update_user(user_id):
    actor = require_admin()
    user = db.session.get(User, user_id)
    if not user or not user.account:
        abort(404, "账号不存在")
    data = _json()
    if "is_active" in data:
        if user.id == actor.id and not bool(data["is_active"]):
            abort(400, "不能禁用当前管理员账号")
        user.account.is_active = bool(data["is_active"])
        user.account.session_version = (user.account.session_version or 0) + 1
    if data.get("name"):
        user.name = data["name"].strip()
    add_audit(
        actor,
        "admin.user_updated",
        "user",
        user.id,
        {"is_active": user.account.is_active},
    )
    db.session.commit()
    return jsonify({"user": user.to_dict()})


@admin_bp.post("/users/<int:user_id>/reset-password")
def reset_password(user_id):
    actor = require_admin()
    user = db.session.get(User, user_id)
    if not user or not user.account:
        abort(404, "账号不存在")
    password = _json().get("password") or _temporary_password()
    _validate_password(password)
    user.account.set_password(password, must_change=True)
    add_audit(actor, "admin.password_reset", "user", user.id)
    db.session.commit()
    return jsonify({"user": user.to_dict(), "initial_password": password})


@admin_bp.post("/terms")
def create_term():
    actor = require_admin()
    data = _json()
    code = (data.get("code") or "").strip()
    name = (data.get("name") or code).strip()
    if not code:
        abort(400, "学期代码不能为空")
    if AcademicTerm.query.filter_by(code=code).first():
        abort(400, "学期代码已存在")
    item = AcademicTerm(code=code, name=name)
    db.session.add(item)
    db.session.flush()
    add_audit(actor, "admin.term_created", "academic_term", item.id)
    db.session.commit()
    return jsonify({"term": item.to_dict()}), 201


@admin_bp.post("/classes")
def create_class():
    actor = require_admin()
    data = _json()
    term = db.session.get(AcademicTerm, data.get("term_id"))
    code = (data.get("code") or "").strip()
    name = (data.get("name") or code).strip()
    if not term or not code:
        abort(400, "学期和班级代码不能为空")
    item = ClassGroup(term_id=term.id, code=code, name=name)
    db.session.add(item)
    db.session.flush()
    add_audit(actor, "admin.class_created", "class_group", item.id)
    db.session.commit()
    return jsonify({"class_group": item.to_dict()}), 201


@admin_bp.post("/courses")
def create_course():
    actor = require_admin()
    data = _json()
    term = db.session.get(AcademicTerm, data.get("term_id"))
    code = (data.get("code") or "").strip()
    name = (data.get("name") or code).strip()
    if not term or not code:
        abort(400, "学期和课程代码不能为空")
    item = Course(term_id=term.id, code=code, name=name)
    db.session.add(item)
    db.session.flush()
    add_audit(actor, "admin.course_created", "course", item.id)
    db.session.commit()
    return jsonify({"course": item.to_dict()}), 201


@admin_bp.post("/enrollments")
def create_enrollment():
    actor = require_admin()
    data = _json()
    student = db.session.get(User, data.get("student_id"))
    term, class_group, course = _academic_items(
        data.get("term_id"), data.get("class_id"), data.get("course_id")
    )
    if not student or student.role != "student":
        abort(400, "学生账号不存在")
    item = CourseEnrollment.query.filter_by(
        student_id=student.id, class_id=class_group.id, course_id=course.id
    ).first()
    if not item:
        item = CourseEnrollment(
            student_id=student.id, class_id=class_group.id, course_id=course.id
        )
        db.session.add(item)
    item.is_active = True
    add_audit(actor, "admin.enrollment_saved", "user", student.id)
    db.session.commit()
    return jsonify({"enrollment": item.to_dict()}), 201


@admin_bp.post("/teaching-assignments")
def create_assignment():
    actor = require_admin()
    data = _json()
    teacher = db.session.get(User, data.get("teacher_id"))
    term, class_group, course = _academic_items(
        data.get("term_id"), data.get("class_id"), data.get("course_id")
    )
    if not teacher or teacher.role != "teacher":
        abort(400, "教师账号不存在")
    item = TeachingAssignment.query.filter_by(
        teacher_id=teacher.id, class_id=class_group.id, course_id=course.id
    ).first()
    if not item:
        item = TeachingAssignment(
            teacher_id=teacher.id, class_id=class_group.id, course_id=course.id
        )
        db.session.add(item)
    item.is_active = True
    add_audit(actor, "admin.assignment_saved", "user", teacher.id)
    db.session.commit()
    return jsonify({"assignment": item.to_dict()}), 201


@admin_bp.patch("/teaching-assignments/<int:assignment_id>")
def update_assignment(assignment_id):
    actor = require_admin()
    item = db.session.get(TeachingAssignment, assignment_id)
    if not item:
        abort(404, "授课关系不存在")
    item.is_active = bool(_json().get("is_active", item.is_active))
    add_audit(
        actor,
        "admin.assignment_updated",
        "teaching_assignment",
        item.id,
        {"is_active": item.is_active},
    )
    db.session.commit()
    return jsonify({"assignment": item.to_dict()})


@admin_bp.patch("/enrollments/<int:enrollment_id>")
def update_enrollment(enrollment_id):
    actor = require_admin()
    item = db.session.get(CourseEnrollment, enrollment_id)
    if not item:
        abort(404, "选课关系不存在")
    item.is_active = bool(_json().get("is_active", item.is_active))
    add_audit(
        actor,
        "admin.enrollment_updated",
        "course_enrollment",
        item.id,
        {"is_active": item.is_active},
    )
    db.session.commit()
    return jsonify({"enrollment": item.to_dict()})


def _update_academic_status(model, item_id, target_type):
    actor = require_admin()
    item = db.session.get(model, item_id)
    if not item:
        abort(404, "教学组织记录不存在")
    item.is_active = bool(_json().get("is_active", item.is_active))
    add_audit(
        actor,
        f"admin.{target_type}_updated",
        target_type,
        item.id,
        {"is_active": item.is_active},
    )
    db.session.commit()
    return jsonify({target_type: item.to_dict()})


@admin_bp.patch("/terms/<int:item_id>")
def update_term(item_id):
    return _update_academic_status(AcademicTerm, item_id, "term")


@admin_bp.patch("/classes/<int:item_id>")
def update_class(item_id):
    return _update_academic_status(ClassGroup, item_id, "class_group")


@admin_bp.patch("/courses/<int:item_id>")
def update_course(item_id):
    return _update_academic_status(Course, item_id, "course")


@admin_bp.post("/students/import")
def import_students():
    actor = require_admin()
    upload = request.files.get("file")
    if not upload:
        abort(400, "请选择学生名单文件")
    rows = _read_import_rows(upload)
    if not rows:
        abort(400, "名单中没有可导入的数据")
    if len(rows) > 5000:
        abort(400, "单次最多导入 5000 名学生")

    results = []
    errors = []
    for index, row in enumerate(rows, start=2):
        required = ["student_no", "name", "class_code", "course_code", "term_code"]
        missing = [field for field in required if not row[field]]
        if missing:
            errors.append({"row": index, "error": f"缺少字段：{', '.join(missing)}"})
            continue
        try:
            if row["initial_password"] and len(row["initial_password"]) < current_app.config[
                "PASSWORD_MIN_LENGTH"
            ]:
                raise ValueError(
                    f"初始密码至少需要 {current_app.config['PASSWORD_MIN_LENGTH']} 位"
                )
            user = User.query.filter_by(student_no=row["student_no"]).first()
            if user and user.role != "student":
                raise ValueError("人员编号已被非学生账号使用")
            login_account = UserAccount.query.filter_by(login_id=row["student_no"]).first()
            if login_account and (not user or login_account.user_id != user.id):
                raise ValueError("登录账号已被其他用户使用")

            term = AcademicTerm.query.filter_by(code=row["term_code"]).first()
            if not term:
                term = AcademicTerm(code=row["term_code"], name=row["term_name"])
                db.session.add(term)
                db.session.flush()
            class_group = ClassGroup.query.filter_by(
                term_id=term.id, code=row["class_code"]
            ).first()
            if not class_group:
                class_group = ClassGroup(
                    term_id=term.id, code=row["class_code"], name=row["class_name"]
                )
                db.session.add(class_group)
                db.session.flush()
            course = Course.query.filter_by(
                term_id=term.id, code=row["course_code"]
            ).first()
            if not course:
                course = Course(
                    term_id=term.id, code=row["course_code"], name=row["course_name"]
                )
                db.session.add(course)
                db.session.flush()

            if not user:
                user = User(student_no=row["student_no"], name=row["name"], role="student")
                db.session.add(user)
                db.session.flush()
            else:
                user.name = row["name"]

            initial_password = ""
            if not user.account:
                initial_password = row["initial_password"] or _temporary_password()
                account = UserAccount(user_id=user.id, login_id=row["student_no"])
                account.set_password(initial_password, must_change=True)
                db.session.add(account)

            enrollment = CourseEnrollment.query.filter_by(
                student_id=user.id,
                class_id=class_group.id,
                course_id=course.id,
            ).first()
            if not enrollment:
                enrollment = CourseEnrollment(
                    student_id=user.id,
                    class_id=class_group.id,
                    course_id=course.id,
                )
                db.session.add(enrollment)
            enrollment.is_active = True
            results.append(
                {
                    "student_no": user.student_no,
                    "name": user.name,
                    "class_name": class_group.name,
                    "course_name": course.name,
                    "term_name": term.name,
                    "initial_password": initial_password,
                }
            )
        except ValueError as exc:
            errors.append({"row": index, "error": str(exc)})

    add_audit(
        actor,
        "admin.students_imported",
        "student_import",
        None,
        {"imported": len(results), "errors": len(errors)},
    )
    db.session.commit()
    return jsonify({"imported": results, "errors": errors, "count": len(results)})
